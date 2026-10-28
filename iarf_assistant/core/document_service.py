from __future__ import annotations

import hashlib
import json
import mimetypes
import re
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pypdfium2 as pdfium
from pypdf import PdfReader, PdfWriter

from iarf_assistant.core.runtime_contracts import bounded_text, estimate_tokens
from iarf_assistant.core.runtime_state import (
    create_local_file_artifact,
    resolve_artifact_path,
    visible_to_user,
)
from iarf_assistant.core.upload_paths import resolve_upload_reference, uploads_root
from carlina.core.utc import utc_now_db
from iarf_assistant.extensions import db
from iarf_assistant.models import RuntimeArtifact, RuntimeResourceReference

DOCUMENT_DIRECT_SUFFIXES = {
    ".csv",
    ".docx",
    ".json",
    ".md",
    ".pdf",
    ".txt",
}
DOCUMENT_ADVANCED_SUFFIXES = {
    ".doc",
    ".epub",
    ".fb2",
    ".htm",
    ".html",
    ".msg",
    ".odt",
    ".rtf",
}
DOCUMENT_PREVIEW_MAX_PAGES = 12
DOCUMENT_EXTRACT_MAX_PAGES = 200
DOCUMENT_SPLIT_MAX_PAGES = 100


@dataclass(frozen=True)
class DocumentResource:
    reference: str
    path: Path
    label: str
    source_kind: str
    artifact_id: str | None = None
    resource_reference_id: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "reference": self.reference,
            "label": self.label,
            "source_kind": self.source_kind,
            "artifact_id": self.artifact_id,
            "resource_reference_id": self.resource_reference_id,
            "suffix": self.path.suffix.lower(),
        }


class DocumentService:
    """Owner-scoped, page-aware document operations backed by runtime artifacts."""

    def __init__(
        self,
        *,
        instance_path: str | Path,
        app_root_path: str | Path,
        owner_user_id: int | None,
        is_admin: bool = False,
        workflow_run_id: str | None = None,
        workflow_step_id: str | None = None,
        advanced_extractor: Callable[..., Any] | None = None,
    ):
        self.instance_path = Path(instance_path).resolve()
        self.app_root_path = Path(app_root_path).resolve()
        self.owner_user_id = owner_user_id
        self.is_admin = bool(is_admin)
        self.workflow_run_id = workflow_run_id
        self.workflow_step_id = workflow_step_id
        self.advanced_extractor = advanced_extractor

    def inspect(self, file_ref: Any) -> dict[str, Any]:
        resource = self.resolve(file_ref)
        path = resource.path
        stat = path.stat()
        mime_type, _ = mimetypes.guess_type(str(path))
        content_hash = _file_sha256(path)
        suffix = path.suffix.lower()
        result: dict[str, Any] = {
            "status": "ok",
            "resource": resource.to_dict(),
            "filename": path.name,
            "mime_type": mime_type or "application/octet-stream",
            "size_bytes": stat.st_size,
            "content_hash": content_hash,
            "direct_operations": suffix in DOCUMENT_DIRECT_SUFFIXES,
            "advanced_provider_available": callable(self.advanced_extractor),
        }

        if suffix == ".pdf":
            document = PdfReader(path, strict=False)
            encrypted = bool(document.is_encrypted)
            operations = ["inspect"]
            if not encrypted:
                operations.extend([
                    "preview_pages",
                    "extract_pages",
                    "split_pages",
                ])
            result.update({
                "page_count": len(document.pages),
                "metadata": _compact_pdf_metadata(document.metadata),
                "encrypted": encrypted,
                "needs_password": encrypted,
                "available_operations": operations,
            })
        elif suffix == ".docx":
            try:
                from docx import Document

                document = Document(path)
                result["paragraph_count"] = len(document.paragraphs)
                result["table_count"] = len(document.tables)
                result["available_operations"] = ["inspect", "extract"]
            except ImportError:
                result["available_operations"] = ["inspect", "advanced_extract"]
                result["warning"] = "python-docx is unavailable; use the advanced document provider."
        elif suffix in {".txt", ".md", ".csv", ".json"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            result["character_count"] = len(text)
            result["estimated_tokens"] = estimate_tokens(text)
            result["available_operations"] = ["inspect", "extract"]
        else:
            result["available_operations"] = ["inspect", "advanced_extract"]
            if suffix not in DOCUMENT_ADVANCED_SUFFIXES:
                result["warning"] = f"Document type '{suffix or 'unknown'}' needs an installed advanced provider."
        return result

    def preview_pages(
        self,
        file_ref: Any,
        *,
        pages: Any = None,
        dpi: int = 120,
        max_pages: int = DOCUMENT_PREVIEW_MAX_PAGES,
    ) -> dict[str, Any]:
        resource = self.resolve(file_ref)
        if resource.path.suffix.lower() != ".pdf":
            raise ValueError("Page preview currently supports PDF resources.")
        render_dpi = max(72, min(int(dpi or 120), 220))
        operation_id = uuid.uuid4().hex[:16]
        content_hash = _file_sha256(resource.path)
        output_dir = (
            self.instance_path
            / "artifacts"
            / "documents"
            / content_hash[:20]
            / "previews"
            / operation_id
        )
        output_dir.mkdir(parents=True, exist_ok=True)

        artifacts: list[dict[str, Any]] = []
        document = pdfium.PdfDocument(str(resource.path))
        try:
            page_count = len(document)
            selected = parse_page_selection(
                pages,
                total_pages=page_count,
                default_pages=[1],
                max_pages=min(DOCUMENT_PREVIEW_MAX_PAGES, max(1, int(max_pages or 1))),
            )
            for page_number in selected:
                page = document.get_page(page_number - 1)
                try:
                    target = output_dir / f"page-{page_number:04d}.png"
                    bitmap = page.render(scale=render_dpi / 72.0)
                    try:
                        image = bitmap.to_pil().convert("RGB")
                        try:
                            image.save(target, format="PNG")
                            width, height = image.size
                        finally:
                            image.close()
                    finally:
                        bitmap.close()
                finally:
                    page.close()
                artifact = self._create_artifact(
                    target,
                    label=f"{resource.label} — page {page_number} preview",
                    metadata={
                        "kind": "document_page_preview",
                        "source_reference": resource.reference,
                        "source_content_hash": content_hash,
                        "page_number": page_number,
                        "page_count": page_count,
                        "dpi": render_dpi,
                        "renderer": "pypdfium2",
                    },
                    source_ids=[resource.reference],
                )
                artifacts.append({
                    "page_number": page_number,
                    "width": width,
                    "height": height,
                    "artifact": artifact.to_dict(),
                })
        finally:
            document.close()

        return {
            "status": "ok",
            "resource": resource.to_dict(),
            "selected_pages": selected,
            "rendered_page_count": len(artifacts),
            "dpi": render_dpi,
            "previews": artifacts,
            "artifact_ids": [row["artifact"]["id"] for row in artifacts],
        }

    def extract(
        self,
        file_ref: Any,
        *,
        pages: Any = None,
        output_format: str = "markdown",
        engine: str = "direct",
        force_ocr: bool = False,
        max_pages: int = DOCUMENT_EXTRACT_MAX_PAGES,
    ) -> dict[str, Any]:
        resource = self.resolve(file_ref)
        normalized_engine = str(engine or "direct").strip().lower()
        normalized_format = str(output_format or "markdown").strip().lower()
        if normalized_format not in {"markdown", "text", "json"}:
            raise ValueError("output_format must be markdown, text, or json.")

        page_numbers: list[int] = []
        provider = "document_service.direct"
        warnings: list[str] = []
        if normalized_engine in {"advanced", "doc_processor", "ocr", "layout"}:
            if not callable(self.advanced_extractor):
                raise RuntimeError("The advanced document provider is not available.")
            if resource.path.suffix.lower() == ".pdf":
                document = PdfReader(resource.path, strict=False)
                page_count = len(document.pages)
                page_numbers = parse_page_selection(
                    pages,
                    total_pages=page_count,
                    default_pages=list(range(1, page_count + 1)),
                    max_pages=min(
                        DOCUMENT_EXTRACT_MAX_PAGES,
                        max(1, int(max_pages or DOCUMENT_EXTRACT_MAX_PAGES)),
                    ),
                )
                page_range = (
                    _page_selection_string(pages)
                    or _page_selection_string(page_numbers)
                )
            else:
                if pages not in (None, "", [], {}):
                    raise ValueError(
                        "Page selection for advanced extraction is available only for PDF resources."
                    )
                page_range = ""
            extracted = self.advanced_extractor(
                path=resource.path,
                page_range=page_range,
                engine=normalized_engine,
                force_ocr=force_ocr,
            )
            if isinstance(extracted, dict):
                content = str(
                    extracted.get("content")
                    or extracted.get("text")
                    or extracted.get("result")
                    or ""
                )
                provider = str(extracted.get("provider") or "doc_processor")
                warnings.extend(str(item) for item in extracted.get("warnings") or [])
            else:
                content = str(extracted or "")
                provider = "doc_processor"
            if content.startswith("Error:"):
                raise RuntimeError(content)
        else:
            content, page_numbers, direct_warnings = self._extract_direct(
                resource,
                pages=pages,
                output_format=normalized_format,
                max_pages=max_pages,
            )
            warnings.extend(direct_warnings)

        extension = {"markdown": ".md", "text": ".txt", "json": ".json"}[normalized_format]
        operation_id = uuid.uuid4().hex[:16]
        content_hash = _file_sha256(resource.path)
        output_dir = (
            self.instance_path
            / "artifacts"
            / "documents"
            / content_hash[:20]
            / "extractions"
        )
        output_dir.mkdir(parents=True, exist_ok=True)
        target = output_dir / f"extract-{operation_id}{extension}"
        if normalized_format == "json":
            try:
                parsed = json.loads(content)
            except json.JSONDecodeError:
                parsed = {
                    "source": resource.reference,
                    "pages": page_numbers,
                    "text": content,
                }
            target.write_text(
                json.dumps(parsed, indent=2, ensure_ascii=False),
                encoding="utf-8",
            )
            stored_content = target.read_text(encoding="utf-8")
        else:
            target.write_text(content, encoding="utf-8")
            stored_content = content

        artifact = self._create_artifact(
            target,
            label=f"{resource.label} — extracted content",
            metadata={
                "kind": "document_text_extraction",
                "source_reference": resource.reference,
                "source_content_hash": content_hash,
                "selected_pages": page_numbers,
                "provider": provider,
                "engine": normalized_engine,
                "output_format": normalized_format,
                "estimated_tokens": estimate_tokens(stored_content),
            },
            source_ids=[resource.reference],
        )
        return {
            "status": "ok",
            "resource": resource.to_dict(),
            "provider": provider,
            "engine": normalized_engine,
            "output_format": normalized_format,
            "selected_pages": page_numbers,
            "character_count": len(stored_content),
            "estimated_tokens": estimate_tokens(stored_content),
            "content_preview": bounded_text(stored_content, 1600),
            "artifact": artifact.to_dict(),
            "artifact_ids": [artifact.id],
            "data_refs": [{
                "id": f"artifact:{artifact.id}",
                "kind": "artifact",
                "label": artifact.label,
                "mime_type": artifact.mime_type,
                "size_bytes": artifact.size_bytes,
                "estimated_tokens": estimate_tokens(stored_content),
                "preview": bounded_text(stored_content, 600),
                "lineage": {
                    "source_reference": resource.reference,
                    "selected_pages": page_numbers,
                },
            }],
            "warnings": warnings,
        }

    def split_pages(
        self,
        file_ref: Any,
        *,
        pages: Any,
        mode: str = "selection",
        max_pages: int = DOCUMENT_SPLIT_MAX_PAGES,
    ) -> dict[str, Any]:
        resource = self.resolve(file_ref)
        if resource.path.suffix.lower() != ".pdf":
            raise ValueError("Page splitting currently supports PDF resources.")
        normalized_mode = str(mode or "selection").strip().lower()
        if normalized_mode not in {"selection", "individual"}:
            raise ValueError("mode must be selection or individual.")

        operation_id = uuid.uuid4().hex[:16]
        content_hash = _file_sha256(resource.path)
        output_dir = (
            self.instance_path
            / "artifacts"
            / "documents"
            / content_hash[:20]
            / "splits"
            / operation_id
        )
        output_dir.mkdir(parents=True, exist_ok=True)
        outputs: list[dict[str, Any]] = []

        document = PdfReader(resource.path, strict=False)
        selected = parse_page_selection(
            pages,
            total_pages=len(document.pages),
            default_pages=None,
            max_pages=min(DOCUMENT_SPLIT_MAX_PAGES, max(1, int(max_pages or 1))),
        )
        groups = [selected] if normalized_mode == "selection" else [[page] for page in selected]
        for index, group in enumerate(groups, start=1):
            output = PdfWriter()
            for page_number in group:
                output.add_page(document.pages[page_number - 1])
            filename = (
                f"pages-{selected[0]:04d}-{selected[-1]:04d}.pdf"
                if normalized_mode == "selection"
                else f"page-{group[0]:04d}.pdf"
            )
            target = output_dir / filename
            with target.open("wb") as handle:
                output.write(handle)
            artifact = self._create_artifact(
                target,
                label=f"{resource.label} — pages {', '.join(str(page) for page in group)}",
                metadata={
                    "kind": "document_page_split",
                    "source_reference": resource.reference,
                    "source_content_hash": content_hash,
                    "selected_pages": group,
                    "split_mode": normalized_mode,
                    "part_index": index,
                    "part_count": len(groups),
                },
                source_ids=[resource.reference],
            )
            outputs.append({
                "pages": group,
                "artifact": artifact.to_dict(),
            })

        return {
            "status": "ok",
            "resource": resource.to_dict(),
            "mode": normalized_mode,
            "selected_pages": selected,
            "output_count": len(outputs),
            "outputs": outputs,
            "artifact_ids": [row["artifact"]["id"] for row in outputs],
        }

    def resolve(self, file_ref: Any) -> DocumentResource:
        reference = str(file_ref or "").strip()
        if not reference:
            raise ValueError("file_ref is required.")

        artifact_id = reference.split(":", 1)[1] if reference.startswith("artifact:") else reference
        artifact = db.session.get(RuntimeArtifact, artifact_id)
        if artifact is not None:
            if self.owner_user_id is not None and not visible_to_user(
                artifact,
                self.owner_user_id,
                self.is_admin,
            ):
                raise PermissionError("Document artifact is not visible to the current user.")
            if (
                self.workflow_run_id
                and artifact.workflow_run_id
                and artifact.workflow_run_id != self.workflow_run_id
            ):
                raise PermissionError("Document artifact belongs to a different workflow run.")
            path = resolve_artifact_path(artifact, instance_path=self.instance_path)
            if not path or not path.exists() or not path.is_file():
                raise FileNotFoundError("Document artifact file is unavailable.")
            return DocumentResource(
                reference=f"artifact:{artifact.id}",
                path=path,
                label=artifact.label or path.name,
                source_kind="artifact",
                artifact_id=artifact.id,
            )

        resource_id = reference.split(":", 1)[1] if reference.startswith("resource:") else reference
        grant = db.session.get(RuntimeResourceReference, resource_id)
        if grant is not None:
            if (
                self.owner_user_id is not None
                and grant.owner_user_id not in {None, self.owner_user_id}
                and not self.is_admin
            ):
                raise PermissionError("Document resource grant is not visible to the current user.")
            if (
                self.workflow_run_id
                and grant.workflow_run_id
                and grant.workflow_run_id != self.workflow_run_id
            ):
                raise PermissionError(
                    "Document resource grant belongs to a different workflow run."
                )
            permissions = {
                str(item).strip().lower()
                for item in (grant.permissions or ["read"])
                if str(item).strip()
            }
            if "read" not in permissions and not self.is_admin:
                raise PermissionError("Document resource grant does not allow reading.")
            if grant.expires_at is not None:
                if grant.expires_at <= utc_now_db():
                    raise PermissionError("Document resource grant has expired.")
            path = self._resolve_granted_uri(grant.resource_uri)
            return DocumentResource(
                reference=f"resource:{grant.id}",
                path=path,
                label=grant.label or path.name,
                source_kind="resource_reference",
                resource_reference_id=grant.id,
            )

        path = self._resolve_upload(reference)
        return DocumentResource(
            reference=reference,
            path=path,
            label=path.name,
            source_kind="upload",
        )

    def _resolve_upload(self, reference: str) -> Path:
        path = resolve_upload_reference(
            reference,
            self.app_root_path,
            require_within_uploads=True,
        )
        if path is None or not path.exists() or not path.is_file():
            raise FileNotFoundError("Document upload is unavailable.")
        root = uploads_root(self.app_root_path)
        relative = path.relative_to(root)
        if (
            self.owner_user_id is not None
            and not self.is_admin
            and (not relative.parts or relative.parts[0] != str(self.owner_user_id))
        ):
            raise PermissionError("Document upload is not owned by the current user.")
        return path

    def _resolve_granted_uri(self, uri: str) -> Path:
        value = str(uri or "").strip()
        if value.startswith("artifact:"):
            return self.resolve(value).path
        if value.startswith("instance://"):
            relative = value[len("instance://"):].lstrip("/\\")
            path = (self.instance_path / relative).resolve()
            path.relative_to(self.instance_path)
        elif value.startswith("/uploads/") or value.startswith("uploads/"):
            path = self._resolve_upload(value)
        else:
            path = Path(value).expanduser().resolve()
        if not path.exists() or not path.is_file():
            raise FileNotFoundError("Granted document resource is unavailable.")
        return path

    def _extract_direct(
        self,
        resource: DocumentResource,
        *,
        pages: Any,
        output_format: str,
        max_pages: int,
    ) -> tuple[str, list[int], list[str]]:
        suffix = resource.path.suffix.lower()
        warnings: list[str] = []
        if suffix == ".pdf":
            document = PdfReader(resource.path, strict=False)
            page_count = len(document.pages)
            selected = parse_page_selection(
                pages,
                total_pages=page_count,
                default_pages=list(range(1, page_count + 1)),
                max_pages=min(
                    DOCUMENT_EXTRACT_MAX_PAGES,
                    max(1, int(max_pages or DOCUMENT_EXTRACT_MAX_PAGES)),
                ),
            )
            rows = []
            for page_number in selected:
                text = (document.pages[page_number - 1].extract_text() or "").strip()
                rows.append({
                    "page": page_number,
                    "text": text,
                })
            if sum(len(row["text"]) for row in rows) < max(20, len(rows) * 8):
                warnings.append(
                    "Very little embedded text was found; retry with engine='advanced' and force_ocr=true."
                )
            if output_format == "json":
                content = json.dumps(
                    {
                        "source": resource.reference,
                        "pages": rows,
                    },
                    indent=2,
                    ensure_ascii=False,
                )
            else:
                marker = "\n\n" if output_format == "text" else "\n\n---\n\n"
                content = marker.join(
                    (
                        f"Page {row['page']}\n\n{row['text']}"
                        if output_format == "text"
                        else f"## Page {row['page']}\n\n{row['text']}"
                    )
                    for row in rows
                )
            return content, selected, warnings

        if pages not in (None, "", [], {}):
            raise ValueError("Page selection is available only for PDF resources.")
        if suffix in {".txt", ".md", ".csv", ".json"}:
            content = resource.path.read_text(encoding="utf-8", errors="replace")
            if output_format == "json" and suffix != ".json":
                content = json.dumps(
                    {"source": resource.reference, "text": content},
                    indent=2,
                    ensure_ascii=False,
                )
            return content, [], warnings
        if suffix == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError(
                    "python-docx is unavailable; use engine='advanced'."
                ) from exc
            document = Document(resource.path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            content = "\n\n".join(text for text in paragraphs if text.strip())
            if output_format == "json":
                content = json.dumps(
                    {
                        "source": resource.reference,
                        "paragraphs": paragraphs,
                    },
                    indent=2,
                    ensure_ascii=False,
                )
            return content, [], warnings
        raise ValueError(
            f"Direct extraction does not support '{suffix or 'unknown'}'; use engine='advanced'."
        )

    def _create_artifact(
        self,
        path: Path,
        *,
        label: str,
        metadata: dict[str, Any],
        source_ids: Any = None,
    ) -> RuntimeArtifact:
        return create_local_file_artifact(
            instance_path=self.instance_path,
            file_path=path,
            workflow_run_id=self.workflow_run_id,
            step_id=self.workflow_step_id,
            created_by_user_id=self.owner_user_id,
            label=label,
            source_ids=source_ids,
            metadata=metadata,
            change_summary="Created by the shared document service.",
            review_status="draft",
            retention_policy="document_runtime",
            download_permission="owner",
        )


def parse_page_selection(
    pages: Any,
    *,
    total_pages: int,
    default_pages: list[int] | None,
    max_pages: int,
) -> list[int]:
    if total_pages <= 0:
        return []
    selected: list[int] = []
    if pages in (None, "", [], {}):
        if default_pages is None:
            raise ValueError("pages is required.")
        selected = list(default_pages)
    elif isinstance(pages, (list, tuple, set)):
        for item in pages:
            try:
                selected.append(int(item))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Invalid page number: {item}") from exc
    else:
        text = str(pages).strip()
        for token in [part.strip() for part in text.split(",") if part.strip()]:
            match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", token)
            if match:
                start, end = int(match.group(1)), int(match.group(2))
                if end < start:
                    raise ValueError(f"Invalid descending page range: {token}")
                selected.extend(range(start, end + 1))
            elif token.isdigit():
                selected.append(int(token))
            else:
                raise ValueError(f"Invalid page selection token: {token}")

    deduplicated = list(dict.fromkeys(selected))
    invalid = [page for page in deduplicated if page < 1 or page > total_pages]
    if invalid:
        raise ValueError(
            f"Page selection is outside the document range 1-{total_pages}: {invalid}"
        )
    limit = max(1, int(max_pages or 1))
    if len(deduplicated) > limit:
        raise ValueError(
            f"Page selection contains {len(deduplicated)} pages; the operation limit is {limit}."
        )
    return deduplicated


def _page_selection_string(pages: Any) -> str:
    if pages in (None, "", [], {}):
        return ""
    if isinstance(pages, (list, tuple, set)):
        return ",".join(str(item) for item in pages)
    return str(pages)


def _compact_pdf_metadata(metadata: Any) -> dict[str, str]:
    payload = metadata if isinstance(metadata, dict) else {}
    canonical_keys = {
        "author": "author",
        "creationdate": "creationDate",
        "creator": "creator",
        "format": "format",
        "keywords": "keywords",
        "moddate": "modDate",
        "producer": "producer",
        "subject": "subject",
        "title": "title",
    }
    result = {}
    for key, value in payload.items():
        canonical = canonical_keys.get(str(key).lstrip("/").casefold())
        if canonical and value not in (None, ""):
            result[canonical] = str(value)
    return result


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
