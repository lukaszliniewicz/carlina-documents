from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path
from unittest.mock import Mock

import pypandoc
import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches
from flask import Flask
from PIL import Image

from iarf_assistant.addons.doc_processor import archive_safety
from iarf_assistant.addons.doc_processor.parsers import html_epub_parser
from iarf_assistant.addons.doc_processor.parsers.html_epub_parser import HtmlEpubParser
from iarf_assistant.addons.doc_processor.semantic_document import (
    canonicalize_document_html,
)


def _png_bytes() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (3, 2), (17, 31, 47)).save(output, format="PNG")
    return output.getvalue()


def _write_docx(path: Path, image_bytes: bytes | None = None) -> None:
    document = Document()
    document.add_paragraph("Synthetic DOCX fixture")
    if image_bytes is not None:
        document.add_picture(io.BytesIO(image_bytes), width=Inches(0.2))
    document.save(path)


def _owner_document_dir(tmp_path: Path) -> tuple[Path, Path]:
    uploads_dir = tmp_path / "uploads"
    document_dir = uploads_dir / "7" / "doc_processor_docs" / "doc-1"
    document_dir.mkdir(parents=True)
    return uploads_dir, document_dir


def test_docx_image_member_decodes_valid_uri_path_once():
    assert (
        html_epub_parser._docx_image_member(
            "media/diagram%20%C3%A9.png",
            {"word/media/diagram é.png"},
        )
        == "word/media/diagram é.png"
    )


def test_docx_image_is_materialized_and_canonicalized_idempotently(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    uploads_dir, document_dir = _owner_document_dir(tmp_path)
    source = document_dir / "source.docx"
    image_bytes = _png_bytes()
    _write_docx(source, image_bytes)
    app = Flask(__name__)
    app.config["CARLINA_UPLOADS_DIR"] = str(uploads_dir)
    parser = HtmlEpubParser(str(source), 7, None)
    read_limits: list[int] = []
    original_read_zip_member = html_epub_parser.read_zip_member

    def read_zip_member(
        archive: zipfile.ZipFile,
        member_name: str,
        *,
        max_bytes: int = archive_safety.MAX_ARCHIVE_MEMBER_BYTES,
    ) -> bytes:
        read_limits.append(max_bytes)
        return original_read_zip_member(
            archive,
            member_name,
            max_bytes=max_bytes,
        )

    monkeypatch.setattr(html_epub_parser, "read_zip_member", read_zip_member)

    with app.app_context():
        status, html = parser.process()
        canonical = canonicalize_document_html(
            html,
            user_id=7,
            source_file_path=str(source),
            uploads_dir=uploads_dir,
        )
        image = BeautifulSoup(canonical.html, "html.parser").find("img")
        expected_name = f"docx_{hashlib.sha256(image_bytes).hexdigest()}.png"
        expected_reference = f"/uploads/7/doc_processor_docs/doc-1/{expected_name}"
        asset_path = document_dir / expected_name

        assert status == "COMPLETED"
        assert image is not None
        assert image.get("src") == expected_reference
        assert asset_path.read_bytes() == image_bytes
        assert (
            hashlib.sha256(asset_path.read_bytes()).hexdigest()
            == hashlib.sha256(image_bytes).hexdigest()
        )
        assert len(list(document_dir.glob("docx_*"))) == 1

        second_status, second_html = parser.process()
        second = canonicalize_document_html(
            second_html,
            user_id=7,
            source_file_path=str(source),
            uploads_dir=uploads_dir,
        )

    assert second_status == "COMPLETED"
    assert second.html == canonical.html
    assert asset_path.read_bytes() == image_bytes
    assert len(list(document_dir.glob("docx_*"))) == 1
    assert read_limits == [25 * 1024 * 1024, 25 * 1024 * 1024]


def test_docx_external_and_unsafe_image_sources_are_not_opened(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    uploads_dir, document_dir = _owner_document_dir(tmp_path)
    source = document_dir / "source.docx"
    _write_docx(source)
    sources = (
        "https://example.invalid/image.png",
        "file:///etc/passwd",
        "//example.invalid/image.png",
        "../outside.png",
        "media%2fimage.png",
        "media/%GG.png",
        r"media\image.png",
        "media/image.png?download=1",
        "media/image.png#fragment",
    )
    output = "".join(f'<img src="{source}"/>' for source in sources)
    converted: list[dict[str, object]] = []
    read_member = Mock(side_effect=AssertionError("untrusted image source was opened"))
    monkeypatch.setattr(html_epub_parser, "read_zip_member", read_member)

    def convert_file(*args, **kwargs):
        converted.append(kwargs)
        return output

    monkeypatch.setattr(pypandoc, "convert_file", convert_file)
    html = HtmlEpubParser(str(source), 7, None)._pandoc_convert(source)
    soup = BeautifulSoup(html, "html.parser")

    assert [image.get("src") for image in soup.find_all("img")] == [
        *sources[:3],
        *([None] * (len(sources) - 3)),
    ]
    assert converted == [{"to": "html", "format": None, "extra_args": ["--sandbox"]}]
    read_member.assert_not_called()

    canonical = canonicalize_document_html(
        html,
        user_id=7,
        source_file_path=str(source),
        uploads_dir=uploads_dir,
    )
    assert all(
        image.get("src") is None
        for image in BeautifulSoup(canonical.html, "html.parser").find_all("img")
    )
    assert not list(document_dir.glob("docx_*"))


def test_docx_missing_package_image_fails_without_reading_another_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _, document_dir = _owner_document_dir(tmp_path)
    source = document_dir / "source.docx"
    _write_docx(source)
    monkeypatch.setattr(
        pypandoc,
        "convert_file",
        lambda *args, **kwargs: '<img src="media/missing.png"/>',
    )
    read_member = Mock(side_effect=AssertionError("missing image was opened"))
    monkeypatch.setattr(html_epub_parser, "read_zip_member", read_member)

    with pytest.raises(
        archive_safety.ArchiveSafetyError, match="missing from its package"
    ):
        HtmlEpubParser(str(source), 7, None)._pandoc_convert(source)

    read_member.assert_not_called()


def test_docx_unsafe_zip_is_rejected_before_pandoc(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    source = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("../outside.png", b"synthetic")
    monkeypatch.setattr(
        pypandoc,
        "convert_file",
        lambda *args, **kwargs: pytest.fail("Pandoc was called"),
    )
    with pytest.raises(archive_safety.ArchiveSafetyError, match="unsafe member path"):
        HtmlEpubParser(str(source), 7, None)._pandoc_convert(source)


@pytest.mark.parametrize("existing_kind", ["symlink", "collision"])
def test_docx_image_output_rejects_symlink_or_byte_collision(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, existing_kind: str
):
    _, document_dir = _owner_document_dir(tmp_path)
    source = document_dir / "source.docx"
    image_bytes = _png_bytes()
    _write_docx(source, image_bytes)
    expected_name = f"docx_{hashlib.sha256(image_bytes).hexdigest()}.png"
    output_path = document_dir / expected_name
    html_epub_parser_output = '<img src="media/image1.png"/>'
    monkeypatch.setattr(
        pypandoc,
        "convert_file",
        lambda *args, **kwargs: html_epub_parser_output,
    )

    if existing_kind == "symlink":
        outside = tmp_path / "sentinel.bin"
        outside.write_bytes(b"preserve this synthetic file")
        try:
            output_path.symlink_to(outside)
        except (NotImplementedError, OSError) as exc:
            pytest.skip(
                f"symlinks are unavailable in this test environment: {type(exc).__name__}"
            )
        with pytest.raises(
            archive_safety.ArchiveSafetyError, match="unlinked regular file"
        ):
            HtmlEpubParser(str(source), 7, None)._pandoc_convert(source)
        assert output_path.is_symlink()
        assert outside.read_bytes() == b"preserve this synthetic file"
    else:
        output_path.write_bytes(b"different synthetic bytes")
        with pytest.raises(archive_safety.ArchiveSafetyError, match="conflicts"):
            HtmlEpubParser(str(source), 7, None)._pandoc_convert(source)
        assert output_path.read_bytes() == b"different synthetic bytes"
